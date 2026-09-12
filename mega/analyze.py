"""
تحلیل واقعی فایل — هر چیزی که کاربر بدهد را باز می‌کند، می‌خواند، اندازه می‌گیرد،
نمودار می‌کشد و یک گزارش فارسی تحویل می‌دهد. هیچ‌کدام از این کارها «ادا» نیست؛
همه‌ی عددها از خود فایل بیرون می‌آیند.

پشتیبان‌ها: تصویر، CSV/TSV/Excel، JSON، کد (Python/JS/…)، متن، Word، PDF، ZIP،
صدا و ویدیو. اگر pandas/matplotlib نبود، تحلیل ساده‌تر ولی واقعی ادامه پیدا می‌کند.
"""
from __future__ import annotations

import asyncio
import json
import mimetypes
import os
import re
import shutil
import subprocess
from pathlib import Path
from typing import Any, Optional

# ---------------------------------------------------------------- دسته‌بندی فایل
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif", ".tif", ".tiff", ".heic"}
DATA_EXT = {".csv", ".tsv", ".xlsx", ".xls", ".xlsm", ".ods", ".parquet", ".jsonl"}
CODE_EXT = {".py", ".js", ".mjs", ".ts", ".tsx", ".jsx", ".java", ".c", ".cpp", ".cs", ".go",
            ".rs", ".php", ".rb", ".sh", ".bash", ".sql", ".html", ".css", ".swift", ".kt",
            ".dart", ".lua", ".r", ".m", ".pl", ".yml", ".yaml", ".toml", ".ini"}
DOC_EXT = {".docx", ".doc", ".pdf", ".rtf"}
TEXT_EXT = {".txt", ".md", ".log", ".json", ".xml", ".srt", ".vtt"}
ARCHIVE_EXT = {".zip", ".tar", ".gz", ".tgz", ".bz2", ".xz", ".7z", ".rar"}
AUDIO_EXT = {".mp3", ".wav", ".m4a", ".ogg", ".opus", ".flac", ".aac", ".webm"}
VIDEO_EXT = {".mp4", ".mov", ".mkv", ".avi", ".webm", ".m4v"}


def kind_of(path: Path) -> str:
    ext = path.suffix.lower()
    if ext in IMAGE_EXT:
        return "image"
    if ext in DATA_EXT:
        return "data"
    if ext in CODE_EXT:
        return "code"
    if ext in DOC_EXT:
        return "doc"
    if ext in TEXT_EXT:
        return "text"
    if ext in ARCHIVE_EXT:
        return "archive"
    if ext in AUDIO_EXT:
        return "audio"
    if ext in VIDEO_EXT:
        return "video"
    return "binary"


# ---------------------------------------------------------------- ابزارهای کمکی
def _run(cmd: list[str], timeout: int = 60) -> str:
    try:
        p = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout)
        return (p.stdout or "") + (p.stderr or "")
    except Exception as e:  # noqa: BLE001
        return f"[خطای اجرای ابزار: {e}]"


def _jsonable(obj: Any) -> Any:
    """هر مقدار را به شکل قابل‌ارسال در JSON درمی‌آورد (NaN، numpy، Timestamp…)."""
    try:
        import math
        if obj is None or isinstance(obj, (str, bool, int)):
            return obj
        if isinstance(obj, float):
            return None if (math.isnan(obj) or math.isinf(obj)) else obj
        if isinstance(obj, dict):
            return {str(k): _jsonable(v) for k, v in obj.items()}
        if isinstance(obj, (list, tuple, set)):
            return [_jsonable(x) for x in obj]
        # numpy / pandas / datetime
        item = getattr(obj, "item", None)
        if callable(item):
            try:
                return _jsonable(item())
            except Exception:
                pass
        if hasattr(obj, "isoformat"):
            return obj.isoformat()
        if hasattr(obj, "tolist"):
            return _jsonable(obj.tolist())
        return str(obj)
    except Exception:
        return str(obj)


def _human_size(n: int) -> str:
    for unit in ("B", "KB", "MB", "GB"):
        if n < 1024 or unit == "GB":
            return f"{n:.0f} {unit}" if unit == "B" else f"{n / 1:.1f} {unit}".replace(".0 ", " ")
        n /= 1024.0
    return f"{n:.1f} GB"


def _fa(text: str) -> str:
    """متن آماده‌ی رسم در matplotlib (شکل‌دهی + چپ‌به‌راست بصری)."""
    try:
        import arabic_reshaper
        from bidi.algorithm import get_display
        return get_display(arabic_reshaper.reshape(text))
    except Exception:
        return text


def _matplotlib_setup():
    """matplotlib با فونت فارسی."""
    import matplotlib
    matplotlib.use("Agg")
    from matplotlib import font_manager, rcParams
    fonts = Path(__file__).resolve().parent.parent / "assets" / "fonts"
    for f in fonts.glob("*.ttf"):
        try:
            font_manager.fontManager.addfont(str(f))
        except Exception:
            continue
    names = {f.name for f in font_manager.fontManager.ttflist}
    for cand in ("Vazirmatn", "Vazirmatn Bold", "Vazir", "DejaVu Sans"):
        if cand in names:
            rcParams["font.family"] = cand
            break
    rcParams["axes.unicode_minus"] = False
    return matplotlib


# ---------------------------------------------------------------- تصویر
def analyze_image_local(path: Path) -> dict:
    """تحلیل فنی واقعی تصویر: اندازه، نسبت، رنگ‌ها، روشنایی، وضوح، پیشنهاد برش."""
    out: dict[str, Any] = {"file": path.name, "size_bytes": path.stat().st_size}
    try:
        from PIL import Image, ImageStat
        img = Image.open(path)
        out["format"] = img.format
        out["dimensions"] = f"{img.width}×{img.height}"
        out["aspect"] = round(img.width / max(1, img.height), 3)
        out["megapixels"] = round(img.width * img.height / 1e6, 2)
        out["orientation"] = ("افقی" if img.width > img.height * 1.05 else
                              "عمودی" if img.height > img.width * 1.05 else "مربع")
        small = img.convert("RGB").resize((160, 160))
        stat = ImageStat.Stat(small)
        brightness = sum(stat.mean) / 3
        out["brightness"] = round(brightness, 1)
        out["brightness_label"] = ("تاریک" if brightness < 70 else
                                   "نیمه‌روشن" if brightness < 165 else "روشن")
        out["contrast"] = round(sum(stat.stddev) / 3, 1)
        # رنگ‌های غالب
        pal = small.quantize(colors=6).convert("RGB").getcolors(160 * 160) or []
        pal = sorted([(c, rgb) for c, rgb in pal], reverse=True)[:5]
        out["palette"] = [{"hex": "#%02x%02x%02x" % rgb, "share": round(c / (160 * 160) * 100, 1)}
                          for c, rgb in pal]
        # وضوح تقریبی (واریانس لبه‌ها)
        from PIL import ImageFilter
        edges = small.convert("L").filter(ImageFilter.FIND_EDGES)
        ev = ImageStat.Stat(edges).stddev[0]
        out["sharpness"] = round(ev, 1)
        out["sharpness_label"] = "شارپ" if ev > 18 else "نسبتاً نرم" if ev > 9 else "تار/کم‌جزئیات"
        # پیشنهاد ابعاد برش برای نسبت‌های تبلیغاتی
        crops = {}
        for ratio, (rw, rh) in {"9:16": (9, 16), "1:1": (1, 1), "16:9": (16, 9)}.items():
            if img.width / img.height > rw / rh:
                w = int(img.height * rw / rh)
                crops[ratio] = f"{w}×{img.height} (برش از وسط)"
            else:
                h = int(img.width * rh / rw)
                crops[ratio] = f"{img.width}×{h} (برش از وسط)"
        out["crop_for"] = crops
        out["warnings"] = []
        if img.width < 800:
            out["warnings"].append("عرض کمتر از ۸۰۰ پیکسل است؛ برای چاپ/تبلیغ بزرگ کیفیت کم می‌آید.")
        if out["sharpness"] < 9:
            out["warnings"].append("تصویر کمی تار است؛ اگر روی آن متن می‌گذاریم، متن را بزرگ‌تر می‌کنیم.")
        img.close()
    except Exception as e:  # noqa: BLE001
        out["error"] = f"خواندن تصویر ناموفق: {e}"
    return out


# ---------------------------------------------------------------- داده
def analyze_data(path: Path, out_dir: Path) -> dict:
    """تحلیل واقعی داده‌ی جدولی + نمودار."""
    ext = path.suffix.lower()
    try:
        import pandas as pd
    except Exception:
        return {"kind": "data", "error": "pandas نصب نیست؛ فقط اطلاعات پایه خوانده شد",
                "file": path.name, "size_bytes": path.stat().st_size}
    try:
        if ext in {".xlsx", ".xls", ".xlsm", ".ods"}:
            xls = pd.ExcelFile(path)
            sheet = xls.sheet_names[0]
            df = xls.parse(sheet)
            extra = {"sheets": xls.sheet_names}
        elif ext == ".jsonl":
            df = pd.read_json(path, lines=True)
            extra = {}
        elif ext == ".parquet":
            df = pd.read_parquet(path)
            extra = {}
        else:
            sep = "\t" if ext == ".tsv" else ","
            df = None
            for enc in ("utf-8", "utf-8-sig", "latin-1"):
                for s in (sep, ";", ","):
                    try:
                        df = pd.read_csv(path, sep=s, encoding=enc)
                        if df.shape[1] > 1:
                            break
                    except Exception:
                        df = None
                if df is not None and df.shape[1] > 1:
                    break
            df = df if df is not None else pd.read_csv(path, encoding="utf-8", errors="replace")
            extra = {}
    except Exception as e:  # noqa: BLE001
        return {"kind": "data", "error": f"خواندن جدول ناموفق: {e}", "file": path.name}

    num = df.select_dtypes("number")
    facts: dict[str, Any] = {
        "file": path.name, **extra,
        "rows": int(df.shape[0]), "columns": int(df.shape[1]),
        "column_names": [str(c) for c in df.columns[:40]],
        "dtypes": {str(c): str(t) for c, t in list(df.dtypes.items())[:40]},
        "missing": {str(c): int(v) for c, v in df.isna().sum().head(40).items() if v},
        "duplicated_rows": int(df.duplicated().sum()),
        "numeric_columns": [str(c) for c in num.columns[:20]],
        "describe": {str(c): {k: (None if v != v else round(float(v), 3)) for k, v in s.items()}
                     for c, s in list(num.describe().items())[:12]},
        "categorical_top": {},
    }
    for c in df.columns[:12]:
        if df[c].dtype == object and df[c].nunique(dropna=True) <= 30:
            facts["categorical_top"][str(c)] = {str(k): int(v) for k, v in
                                                df[c].value_counts().head(6).items()}
    if len(num.columns) > 1:
        try:
            corr = num.corr(numeric_only=True).abs()
            pairs = []
            cols = list(corr.columns)
            for i in range(len(cols)):
                for j in range(i + 1, len(cols)):
                    v = corr.iloc[i, j]
                    if v == v:
                        pairs.append((round(float(v), 3), str(cols[i]), str(cols[j])))
            pairs.sort(reverse=True)
            facts["top_correlations"] = [{"r": r, "a": a, "b": b} for r, a, b in pairs[:6] if r > 0.35]
        except Exception:
            pass
    facts["head"] = df.head(8).to_dict(orient="records")
    facts["tail"] = df.tail(4).to_dict(orient="records")

    charts: list[str] = []
    try:
        _matplotlib_setup()
        import matplotlib.pyplot as plt
        out_dir.mkdir(parents=True, exist_ok=True)
        # ۱) مقادیر گم‌شده
        miss = df.isna().sum()
        miss = miss[miss > 0].head(12)
        if len(miss):
            fig, ax = plt.subplots(figsize=(7, max(2.4, 0.42 * len(miss))), dpi=130)
            ax.barh([_fa(str(c)) for c in miss.index][::-1], list(miss.values)[::-1], color="#6d6af7")
            ax.set_title(_fa("مقادیر گم‌شده در هر ستون"))
            ax.set_xlabel(_fa("تعداد"))
            fig.tight_layout()
            p = out_dir / "chart_missing.png"
            fig.savefig(p); plt.close(fig); charts.append(str(p))
        # ۲) پرتکرارترین دسته‌ها
        cats = facts["categorical_top"]
        if cats:
            col, value = next(iter(cats.items()))
            fig, ax = plt.subplots(figsize=(7, max(2.2, 0.42 * len(value))), dpi=130)
            ax.barh([_fa(str(k)) for k in list(value)[::-1]], list(value.values())[::-1], color="#22c55e")
            ax.set_title(_fa(f"پرتکرارترین مقادیر «{col}»"))
            fig.tight_layout()
            p = out_dir / "chart_categories.png"
            fig.savefig(p); plt.close(fig); charts.append(str(p))
        # ۳) توزیع یک ستون عددی مهم
        if len(num.columns):
            col = max(num.columns, key=lambda c: float(num[c].nunique(dropna=True)))
            series = num[col].dropna()
            if len(series) > 3:
                fig, ax = plt.subplots(figsize=(7, 3.2), dpi=130)
                ax.hist(series.astype(float), bins=24, color="#f59e0b", edgecolor="white")
                ax.set_title(_fa(f"توزیع ستون عددی «{col}»"))
                ax.set_xlabel(_fa(str(col))); ax.set_ylabel(_fa("تعداد"))
                fig.tight_layout()
                p = out_dir / "chart_distribution.png"
                fig.savefig(p); plt.close(fig); charts.append(str(p))
    except Exception as e:  # noqa: BLE001
        facts["chart_error"] = str(e)

    return {"kind": "data", "facts": facts, "charts": charts}


# ---------------------------------------------------------------- کد
def analyze_code(path: Path) -> dict:
    text = path.read_text(encoding="utf-8", errors="replace")
    lines = text.splitlines()
    facts: dict[str, Any] = {
        "file": path.name, "language": path.suffix.lstrip(".") or "?",
        "lines": len(lines),
        "code_lines": sum(1 for l in lines if l.strip() and not l.strip().startswith("#")),
        "comment_lines": sum(1 for l in lines if l.strip().startswith(("#", "//", "/*", "*"))),
        "blank_lines": sum(1 for l in lines if not l.strip()),
        "longest_line": max((len(l) for l in lines), default=0),
        "todos": [f"{i + 1}: {l.strip()[:120]}" for i, l in
                  enumerate(lines) if re.search(r"\b(TODO|FIXME|HACK|XXX)\b", l)][:10],
        "size_bytes": path.stat().st_size,
    }
    ext = path.suffix.lower()
    if ext == ".py":
        try:
            import ast as _ast
            tree = _ast.parse(text)
            funcs = [n.name for n in _ast.walk(tree) if isinstance(n, (_ast.FunctionDef, _ast.AsyncFunctionDef))]
            classes = [n.name for n in _ast.walk(tree) if isinstance(n, _ast.ClassDef)]
            imports = sorted({(n.module or "").split(".")[0] if isinstance(n, _ast.ImportFrom)
                              else (n.names[0].name.split(".")[0] if isinstance(n, _ast.Import) else "")
                              for n in _ast.walk(tree) if isinstance(n, (_ast.Import, _ast.ImportFrom))})
            facts.update({"syntax": "سالم ✅", "functions": funcs[:40], "function_count": len(funcs),
                          "classes": classes[:20], "imports": [i for i in imports if i][:25],
                          "max_depth": _py_depth(tree)})
        except SyntaxError as e:
            facts.update({"syntax": f"خطای نگارشی ❌ خط {e.lineno}: {e.msg}",
                          "syntax_error": {"line": e.lineno, "msg": e.msg, "text": (e.text or "").strip()}})
        except Exception as e:  # noqa: BLE001
            facts["syntax"] = f"بررسی ناموفق: {e}"
    elif ext in {".js", ".mjs", ".cjs", ".ts", ".jsx", ".tsx"} and shutil.which("node"):
        out = _run(["node", "--check", str(path)], timeout=25)
        facts["syntax"] = "سالم ✅" if not out.strip() else f"خطا ❌ {out.strip()[:400]}"
    facts["first_lines"] = "\n".join(lines[:40])
    return {"kind": "code", "facts": facts, "charts": []}


def _py_depth(tree) -> int:
    """ساده‌ترین سنجه‌ی پیچیدگی: بیشترین عمق تودرتویی."""
    import ast as _ast
    best = 0

    def walk(node, d=0):
        nonlocal best
        best = max(best, d)
        for child in _ast.iter_child_nodes(node):
            walk(child, d + 1 if isinstance(child, (_ast.If, _ast.For, _ast.While, _ast.With,
                                                    _ast.Try, _ast.FunctionDef, _ast.ClassDef)) else d)
    walk(tree)
    return best


# ---------------------------------------------------------------- متن / سند
def analyze_text(path: Path) -> dict:
    text = path.read_text(encoding="utf-8", errors="replace")
    words = re.findall(r"[\w\u0600-\u06FF']+", text)
    lines = text.splitlines()
    stop = {"و", "در", "به", "از", "که", "این", "را", "با", "است", "برای", "the", "a", "of", "to",
            "and", "in", "is", "it", "that", "for", "on", "as", "with"}
    freq: dict[str, int] = {}
    for w in words:
        lw = w.lower()
        if len(lw) > 2 and lw not in stop:
            freq[lw] = freq.get(lw, 0) + 1
    top = sorted(freq.items(), key=lambda kv: kv[1], reverse=True)[:20]
    return {"kind": "text", "charts": [], "facts": {
        "file": path.name, "size_bytes": path.stat().st_size,
        "lines": len(lines), "words": len(words),
        "chars": len(text), "unique_words": len(freq),
        "headings": [l.strip() for l in lines if l.strip().startswith("#")][:20],
        "top_words": top,
        "reading_time_min": max(1, round(len(words) / 200)),
        "preview": text[:1200],
    }}


def analyze_doc(path: Path) -> dict:
    facts: dict[str, Any] = {"file": path.name, "size_bytes": path.stat().st_size}
    text = ""
    if path.suffix.lower() == ".docx":
        try:
            import docx
            d = docx.Document(str(path))
            pars = [p.text for p in d.paragraphs if p.text.strip()]
            facts["paragraphs"] = len(pars)
            facts["tables"] = len(d.tables)
            text = "\n".join(pars)
        except Exception as e:  # noqa: BLE001
            facts["error"] = f"خواندن Word ناموفق: {e}"
    elif path.suffix.lower() == ".pdf":
        try:
            from pypdf import PdfReader  # نوعاً نصب نیست
            r = PdfReader(str(path))
            facts["pages"] = len(r.pages)
            text = "\n".join((p.extract_text() or "") for p in r.pages[:60])
        except Exception:
            # بدون pypdf: استخراج خام متن‌های قابل‌خواندن
            raw = path.read_bytes()
            chunks = re.findall(rb"\(([^()]{3,200})\)", raw[:4_000_000])
            text = " ".join(c.decode("latin-1", "ignore") for c in chunks)
            facts["note"] = "pypdf نصب نیست؛ متن به‌صورت خام استخراج شد (نصب: pip install pypdf)"
    if text:
        words = re.findall(r"\S+", text)
        facts.update({"words": len(words), "chars": len(text),
                      "preview": text[:1500], "reading_time_min": max(1, round(len(words) / 200))})
    return {"kind": "doc", "facts": facts, "charts": []}


def analyze_archive(path: Path) -> dict:
    names: list[str] = []
    try:
        if path.suffix.lower() == ".zip":
            import zipfile
            with zipfile.ZipFile(path) as z:
                names = z.namelist()[:400]
        else:
            out = _run(["tar", "-tf", str(path)], timeout=40)
            names = [l for l in out.splitlines() if l.strip()][:400]
    except Exception as e:  # noqa: BLE001
        return {"kind": "archive", "charts": [], "facts": {"file": path.name, "error": str(e)}}
    ext_count: dict[str, int] = {}
    for n in names:
        e = Path(n).suffix.lower() or "(بدون پسوند)"
        ext_count[e] = ext_count.get(e, 0) + 1
    return {"kind": "archive", "charts": [], "facts": {
        "file": path.name, "entries": len(names),
        "top_extensions": dict(sorted(ext_count.items(), key=lambda kv: kv[1], reverse=True)[:12]),
        "first_entries": names[:40]}}


def analyze_media(path: Path) -> dict:
    ext = path.suffix.lower()
    info = _run(["ffprobe", "-v", "error", "-show_entries",
                 "format=duration,size,bit_rate:stream=codec_type,codec_name,width,height,sample_rate,channels",
                 "-of", "json", str(path)], timeout=60)
    facts: dict[str, Any] = {"file": path.name, "size_bytes": path.stat().st_size}
    try:
        d = json.loads(info)
        fmt = d.get("format") or {}
        facts["duration_sec"] = round(float(fmt.get("duration") or 0), 2)
        facts["bitrate_kbps"] = round(float(fmt.get("bit_rate") or 0) / 1000, 1)
        facts["streams"] = [{"type": s.get("codec_type"), "codec": s.get("codec_name"),
                             "size": (f"{s.get('width')}×{s.get('height')}"
                                      if s.get("codec_type") == "video" else
                                      f"{s.get('channels')}ch @ {s.get('sample_rate')}Hz")}
                            for s in (d.get("streams") or [])]
    except Exception as e:  # noqa: BLE001
        facts["error"] = f"ffprobe خواندن نکرد: {e}"
    return {"kind": "audio" if ext in AUDIO_EXT else "video", "facts": facts, "charts": []}


def analyze_binary(path: Path) -> dict:
    raw = path.read_bytes()[:20000]
    printable = sum(1 for b in raw if 32 <= b < 127 or b in (9, 10, 13))
    return {"kind": "binary", "charts": [], "facts": {
        "file": path.name, "size_bytes": path.stat().st_size,
        "mime": mimetypes.guess_type(path.name)[0] or "نامشخص",
        "printable_ratio": round(printable / max(1, len(raw)), 3),
        "hex_preview": raw[:160].hex(" ")}}


# ---------------------------------------------------------------- گزارش فارسی
_KIND_TITLE = {"image": "🖼 تصویر", "data": "📊 داده‌ی جدولی", "code": "💻 کد",
               "text": "📄 متن", "doc": "📝 سند", "archive": "🗜 آرشیو",
               "audio": "🎧 صدا", "video": "🎬 ویدیو", "binary": "📦 فایل"}


def _facts_md(res: dict) -> str:
    kind = res.get("kind", "?")
    f = res.get("facts") or {}
    out: list[str] = [f"### {_KIND_TITLE.get(kind, kind)} — `{f.get('file', '')}`"]

    if kind == "image":
        out += [f"- ابعاد: **{f.get('dimensions')}** · نسبت {f.get('aspect')} · {f.get('orientation')}",
                f"- فرمت: {f.get('format')} · حجم: {f.get('size_bytes', 0) / 1024:.0f} کیلوبایت",
                f"- روشنایی: {f.get('brightness')} ({f.get('brightness_label')}) · "
                f"کنتراست: {f.get('contrast')} · وضوح: {f.get('sharpness')} ({f.get('sharpness_label')})",
                f"- رنگ‌های غالب: " + "، ".join(f"{c['hex']} ({c['share']}٪)" for c in (f.get("palette") or [])),
                "- برش پیشنهادی — " + " · ".join(f"{k}: {v}" for k, v in (f.get("crop_for") or {}).items())]
    elif kind == "data":
        out += [f"- **{f.get('rows')} سطر × {f.get('columns')} ستون** · ستون‌ها: "
                f"{'، '.join(f.get('column_names', [])[:12])}",
                f"- سطرهای تکراری: {f.get('duplicated_rows')} · "
                f"ستون‌های عددی: {len(f.get('numeric_columns', []))}"]
        if f.get("missing"):
            out.append("- مقادیر گم‌شده: " + "، ".join(f"{k}={v}" for k, v in f["missing"].items()))
        else:
            out.append("- مقادیر گم‌شده: ندارد ✅")
        for c, s in list((f.get("describe") or {}).items())[:6]:
            out.append(f"- `{c}`: میانگین {s.get('mean')} · میانه {s.get('50%')} · "
                       f"کمینه {s.get('min')} · بیشینه {s.get('max')}")
        for c, v in list((f.get("categorical_top") or {}).items())[:3]:
            out.append(f"- `{c}` پرتکرار: " + "، ".join(f"{k} ({n})" for k, n in list(v.items())[:5]))
        for cor in (f.get("top_correlations") or [])[:4]:
            out.append(f"- همبستگی {cor['a']} ↔ {cor['b']}: **{cor['r']}**")
    elif kind == "code":
        out += [f"- زبان: {f.get('language')} · {f.get('lines')} خط "
                f"({f.get('code_lines')} کد، {f.get('comment_lines')} توضیح، {f.get('blank_lines')} خالی)",
                f"- بررسی نگارش: **{f.get('syntax')}**"]
        if f.get("function_count") is not None:
            out.append(f"- توابع: {f.get('function_count')} · کلاس‌ها: {len(f.get('classes', []))} · "
                       f"حداکثر تودرتویی: {f.get('max_depth')}")
        if f.get("imports"):
            out.append("- وابستگی‌ها: " + "، ".join(f["imports"][:15]))
        if f.get("todos"):
            out.append("- کارهای ناتمام: " + " | ".join(f["todos"][:5]))
    elif kind == "text":
        out += [f"- {f.get('lines')} خط · {f.get('words')} کلمه · {f.get('chars')} کاراکتر · "
                f"زمان مطالعه ≈ {f.get('reading_time_min')} دقیقه",
                "- پرتکرارترین کلمات: " + "، ".join(f"{w} ({n})" for w, n in (f.get("top_words") or [])[:12])]
        if f.get("headings"):
            out.append("- سرتیترها: " + " | ".join(f["headings"][:8]))
    elif kind == "doc":
        out += [f"- {f.get('words', '?')} کلمه · {f.get('paragraphs', f.get('pages', '?'))} "
                f"{'پاراگراف' if f.get('paragraphs') is not None else 'صفحه'} · "
                f"جدول‌ها: {f.get('tables', 0)}",
                f"- زمان مطالعه ≈ {f.get('reading_time_min', '?')} دقیقه"]
        if f.get("note"):
            out.append(f"- {f['note']}")
    elif kind == "archive":
        out += [f"- {f.get('entries')} آیتم داخل آرشیو",
                "- پسوندها: " + "، ".join(f"{k}={v}" for k, v in (f.get("top_extensions") or {}).items()),
                "- نمونه‌ها: " + "، ".join(f"`{n}`" for n in (f.get("first_entries") or [])[:8])]
    elif kind in {"audio", "video"}:
        out += [f"- مدت: {f.get('duration_sec')} ثانیه · نرخ بیت: {f.get('bitrate_kbps')} kbps"]
        for s in (f.get("streams") or []):
            out.append(f"- {s['type']}: {s['codec']} ({s['size']})")
    else:
        out += [f"- حجم: {f.get('size_bytes', 0) / 1024:.0f} کیلوبایت · نوع: {f.get('mime')}",
                f"- نسبت بایت‌های خوانا: {f.get('printable_ratio')}"]
    if f.get("warnings"):
        out += ["", "**هشدارها:**"] + [f"- ⚠️ {w}" for w in f["warnings"]]
    if f.get("error"):
        out.append(f"- ❌ {f['error']}")
    return "\n".join(out)


# ---------------------------------------------------------------- ورودی اصلی
async def analyze_file(path: str | Path, question: str = "", out_dir: Optional[Path] = None,
                       settings=None, use_model: bool = True, max_seconds: int = 180) -> dict:
    """فایل را واقعاً تحلیل می‌کند و گزارش + نمودار + بررسی مدل برمی‌گرداند."""
    p = Path(path)
    if not p.is_file():
        return {"ok": False, "error": f"فایلی با این مسیر پیدا نشد: {p}"}
    kind = kind_of(p)
    out_dir = Path(out_dir) if out_dir else p.parent / f"{p.stem}_analysis"
    charts: list[str] = []
    facts: dict[str, Any] = {}
    try:
        if kind == "image":
            from .vision import describe_image
            vres = await asyncio.wait_for(describe_image(p, question=question, settings=settings,
                                                         use_model=use_model),
                                          timeout=max_seconds) if use_model else \
                   {"ok": True, "local": analyze_image_local(p), "model_text": ""}
            facts = vres.get("local") or analyze_image_local(p)
            facts["vision"] = vres.get("model_text") or ""
            facts["vision_model"] = vres.get("model") or ("تحلیل محلی" if not vres.get("model_text") else "")
            base = {"kind": "image", "facts": facts, "charts": []}
            md = _facts_md(base)
            if vres.get("model_text"):
                md += "\n\n**توصیف هوش بینایی:**\n\n" + vres["model_text"]
            if vres.get("hint"):
                md += f"\n\n> {vres['hint']}"
            return _jsonable({"ok": True, "kind": "image", "facts": facts, "charts": [],
                              "report": md, "model_review": vres.get("model_text") or ""})

        if kind == "data":
            res = await asyncio.to_thread(analyze_data, p, out_dir)
        elif kind == "code":
            res = await asyncio.to_thread(analyze_code, p)
        elif kind == "text":
            res = await asyncio.to_thread(analyze_text, p)
        elif kind == "doc":
            res = await asyncio.to_thread(analyze_doc, p)
        elif kind == "archive":
            res = await asyncio.to_thread(analyze_archive, p)
        elif kind in {"audio", "video"}:
            res = await asyncio.to_thread(analyze_media, p)
        else:
            res = await asyncio.to_thread(analyze_binary, p)
        facts, charts = res.get("facts", {}), res.get("charts", [])
    except Exception as e:  # noqa: BLE001
        return {"ok": False, "error": f"{type(e).__name__}: {e}", "kind": kind}

    md = _facts_md({"kind": kind, "facts": facts})
    result = {"ok": True, "kind": kind, "facts": facts, "charts": charts, "report": md,
              "model_review": ""}

    # بررسی/جمع‌بندی با مدل — اگر کلید یا حالت نمایشی فعال باشد
    if use_model and kind in {"data", "code", "text", "doc", "archive", "audio", "video", "binary"}:
        try:
            review = await asyncio.wait_for(_model_review(kind, facts, question, settings),
                                            timeout=max_seconds)
            if review:
                result["model_review"] = review
                md += "\n\n**تحلیل هوش‌ها:**\n\n" + review
        except Exception:
            pass
    result["report"] = md
    return _jsonable(result)


async def _model_review(kind: str, facts: dict, question: str, settings) -> str:
    """جمع‌بندی و پاسخ به سؤال کاربر روی یافته‌های واقعی."""
    from .providers import chat, resolve_role
    spec = await resolve_role("judge" if kind in {"data", "code"} else "expert",
                              temperature=0.25, max_tokens=1400, settings=settings)
    if not spec:
        return ""
    compact = {k: v for k, v in facts.items() if k not in ("hex_preview",)}
    sys_msg = ("تو یک تحلیل‌گر فنی دقیق هستی. فقط بر اساس «یافته‌های واقعی» زیر حرف بزن؛ "
               "چیزی از خودت اضافه نکن. فارسی، جمع‌بندی‌شده و کاربردی بنویس: "
               "۱) این فایل دقیقاً چه چیزی است؟ ۲) مهم‌ترین نکته‌ها/مشکلات ۳) پیشنهاد گام بعدی.")
    user = (f"نوع فایل: {kind}\n" + (f"سؤال کاربر: {question}\n" if question else "") +
            "یافته‌ها (JSON):\n" + json.dumps(compact, ensure_ascii=False)[:6000])
    res = await chat(spec, [{"role": "system", "content": sys_msg}, {"role": "user", "content": user}],
                     settings=settings)
    return res.text if res.ok else ""
